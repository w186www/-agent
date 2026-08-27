"""知识库灌入（kb_ingest）测试：清洗 / 切分 / 失败降级 / 成功落库与重传覆盖。

- 清洗与切分为纯函数，直接断言；
- ingest_bidding_doc 成功路径用 fake Embedding / Qdrant（避免网络依赖），
  失败路径（文件不存在）验证 status=2 降级不抛异常。
"""

from unittest.mock import MagicMock

from sqlalchemy import select

from app.ai import kb_ingest
from app.ai.kb_ingest import chunk_text, clean_text
from app.db.models import KbChunk, KbDocument


# ---------------------------------------------------------------------------
# 清洗
# ---------------------------------------------------------------------------


def test_clean_text():
    raw = "第一章\u200b 简介\n\n\n12\n第 二 章\t说明　内容"
    cleaned = clean_text(raw)
    assert cleaned == "第一章 简介\n第 二 章 说明 内容"


def test_clean_text_removes_page_number_only_lines():
    raw = "正文第一行\n  3 \n正文第二行"
    cleaned = clean_text(raw)
    assert "3" not in cleaned


# ---------------------------------------------------------------------------
# 切分
# ---------------------------------------------------------------------------


def test_chunk_text_sentence_fallback():
    """窗口应回退到句子终止符切分，不在句子中间截断（末块除外）。"""
    text = "第一句。" + "第二句很长很长。" * 30
    chunks = chunk_text(text, size=30, overlap=5)
    assert len(chunks) > 1
    for c in chunks:
        assert c.strip().endswith(("。", "！", "？"))


def test_chunk_text_window_size():
    text = "等保测评。" * 300  # 1500 字符
    chunks = chunk_text(text, size=100, overlap=20)
    assert len(chunks) > 1
    assert all(len(c) <= 100 for c in chunks)


def test_chunk_text_md5_dedup():
    text = "重复内容。" * 200
    chunks = chunk_text(text, size=50, overlap=10)
    assert chunks
    assert len(chunks) == len(set(chunks))  # 无重复片段


def test_chunk_text_empty():
    assert chunk_text("   ") == []


def test_extract_docx_heading_markers(tmp_path):
    """DOCX 提取时识别 Heading 1/2/3，注入 # / ## / ### 结构标记。"""
    from docx import Document

    p = tmp_path / "structure.docx"
    docx = Document()
    docx.add_heading("第一章 总体要求", level=1)
    docx.add_paragraph("正文内容")
    docx.add_heading("1.1 概述", level=2)
    docx.add_heading("小节", level=3)
    docx.save(str(p))

    text = kb_ingest.extract_text(str(p))
    assert "# 第一章 总体要求" in text
    assert "## 1.1 概述" in text
    assert "### 小节" in text
    assert "正文内容" in text


# ---------------------------------------------------------------------------
# ingest_bidding_doc
# ---------------------------------------------------------------------------


class FakeEmbeddings:
    """假 Embedding：固定 2 维向量，不触发网络调用。"""

    def embed(self, text: str) -> list[float]:
        return [0.1, 0.2]

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[0.1, 0.2]] * len(texts)


class FakeQdrant:
    """假 Qdrant 客户端：记录 delete/upsert 调用，集合始终存在。"""

    def __init__(self):
        self.delete = MagicMock()
        self.upsert = MagicMock()

    def collection_exists(self, name: str) -> bool:
        return True

    def get_collection(self, name: str):
        return MagicMock(config=MagicMock(params=MagicMock(vectors=MagicMock(size=2))))


def _make_doc(db_session, title="历史标书A") -> KbDocument:
    doc = KbDocument(user_id=1, kb_type=2, doc_title=title, file_path="", chunk_count=0, status=0)
    db_session.add(doc)
    db_session.commit()
    return doc


def test_ingest_failure_marks_status2(client, db_session, tmp_path):
    """文件不存在时返回 status=2 且不抛异常。"""
    doc = _make_doc(db_session)
    result = kb_ingest.ingest_bidding_doc(
        db_session, doc.id, str(tmp_path / "not_exist.pdf"), doc.doc_title, kb_type=2
    )
    assert result["status"] == 2
    assert "文件不存在" in result["error"]
    db_session.refresh(doc)
    assert doc.status == 2


def test_ingest_success_saves_chunks_and_rewrite(client, db_session, monkeypatch, tmp_path):
    """成功路径：向量化入库、原文落 kb_chunks、同文档重传覆盖旧切片。"""
    monkeypatch.setattr(kb_ingest, "get_embeddings", lambda: FakeEmbeddings())
    fake_qdrant = FakeQdrant()
    monkeypatch.setattr(kb_ingest, "get_qdrant_client", lambda: fake_qdrant)

    txt = tmp_path / "bid.txt"
    txt.write_text("第一章 商务要求。\n" + "报价说明：\n" * 120 + "废标项：缺少资质证明。", encoding="utf-8")

    doc = _make_doc(db_session)

    # 首次入库
    result1 = kb_ingest.ingest_bidding_doc(db_session, doc.id, str(txt), doc.doc_title, kb_type=2)
    assert result1["status"] == 1
    assert result1["chunks"] > 0
    db_session.refresh(doc)
    assert doc.status == 1
    assert doc.chunk_count == result1["chunks"]
    assert fake_qdrant.upsert.called

    # 原文落库：vector_id 为 UUID，与 Qdrant point id 一一对应；Qdrant 不存原文
    saved = db_session.scalars(
        select(KbChunk).where(KbChunk.doc_id == doc.id).order_by(KbChunk.chunk_index)
    ).all()
    assert len(saved) == result1["chunks"]
    assert all(c.content for c in saved)
    assert all(len(c.vector_id) == 36 and c.vector_id.count("-") == 4 for c in saved)  # uuid4
    points = fake_qdrant.upsert.call_args.kwargs["points"]
    assert {p.id for p in points} == {c.vector_id for c in saved}
    assert all("text" not in (p.payload or {}) for p in points)  # Qdrant 不存原文
    assert all(p.payload.get("doc_id") == doc.id for p in points)

    # 同文档重传：触发覆盖删除，切片数不累积
    result2 = kb_ingest.ingest_bidding_doc(db_session, doc.id, str(txt), doc.doc_title, kb_type=2)
    assert result2["status"] == 1
    assert fake_qdrant.delete.called
    saved2 = db_session.scalars(
        select(KbChunk).where(KbChunk.doc_id == doc.id)
    ).all()
    assert len(saved2) == result2["chunks"] == result1["chunks"]
