"""Agent 工具层：供流式聊天链路调用的最小可执行工具集。

- search_tender(keyword, max_results)：联网搜索招标公告（Tavily API），返回结构化 JSON 字符串，
  供 LLM 自主调用并筛选推荐；
- download_tender(url)：从招标公告页面提取附件下载链接并下载到本地（tenders 目录），
  返回本地文件路径（搜索→下载→解析三步链式的第二步）；
- parse_tender(file_path)：解析招标文件（.pdf/.docx），提取预算、评分表、资格要求、废标项，
  内部调用 DeepSeek 做结构化提取（三步链式的第三步）；
- search_tender_local(keywords, db)：本地历史标书库检索（Qdrant kb_type=2，RAG 参考用），
  返回相似标书片段与引用 doc_id 列表；
- analyze_screenshot(screenshot_path, checklist_code, checklist_name)：测评截图分析，
  调用 Qwen-VL 多模态模型识别控制点相关配置信息（无 Key / 调用失败时如实返回错误，不伪造结果）。

工具只做单次执行并返回结构化结果，不做链路组装、不定义 Prompt。
"""

import json
import os
import time
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

from langchain_core.tools import tool
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..db.models import AssessmentRecord, KbDocument, ReviewTask
from ..db.session import SessionLocal
from ..log import get_logger
from .llm import get_deepseek, get_qwen_text, get_qwen_vl
from .memory import get_rag_context

logger = get_logger("tools")

# download_tender 请求招标公告页的浏览器 User-Agent（部分站点对默认 UA 返回 403）
_TENDER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)
# 招标文件附件扩展名白名单
_TENDER_FILE_EXTS = (".pdf", ".doc", ".docx", ".zip")


def _parse_llm_json(raw: str) -> dict:
    """清理 LLM 输出中可能包裹的 Markdown 代码块，解析为 dict。

    raw 可能形如 ```json\n{...}\n```，统一剥离围栏后 json.loads。
    """
    text = (raw or "").strip()
    if text.startswith("```"):
        text = text.strip("`").strip()
        if text.lower().startswith("json"):
            text = text[4:].lstrip()
    parsed = json.loads(text)
    if not isinstance(parsed, dict):
        raise ValueError(f"LLM 输出不是 JSON 对象: {text[:100]}")
    return parsed


# 拓扑规范化：设备类型 / 网络层级的兜底默认（LLM 漏填或乱填时按名称推断）
_TOPOLOGY_TYPES = {"switch", "router", "firewall", "server", "database"}
_TOPOLOGY_LAYERS = {"core", "aggregation", "access"}
_TOPOLOGY_TYPE_LAYER = {
    "router": "core",
    "firewall": "core",
    "switch": "aggregation",
    "server": "access",
    "database": "access",
}


def _infer_device_type(label: str) -> str:
    """按设备名称关键词推断设备类型（LLM 未填/乱填 type 时兜底）。"""
    text = label or ""
    if any(k in text for k in ("防火墙", "fw", "FW")):
        return "firewall"
    if any(k in text for k in ("路由", "route")):
        return "router"
    if any(k in text for k in ("数据库", "db", "mysql", "oracle", "redis")):
        return "database"
    if any(k in text for k in ("服务器", "server", "web", "应用", "主机", "虚拟")):
        return "server"
    return "switch"


def _normalize_topology(data: dict) -> dict:
    """规范化 LLM 生成的拓扑数据：补全节点字段、过滤无效边、兜底安全域与摘要。

    DeepSeek 输出结构不稳定（缺 id/type/layer、引用不存在的节点），
    逐层校验补全，避免前端渲染出悬空连线、未知图标、空安全域。
    """
    # 1. 节点：补全字段、去重 id、type/layer 按名称兜底推断
    nodes: list[dict] = []
    used_ids: set[str] = set()
    for index, node in enumerate(data.get("nodes") or [], start=1):
        if not isinstance(node, dict):
            continue
        nid = str(node.get("id") or "").strip() or f"node-{index}"
        while nid in used_ids:
            nid = f"{nid}-{index}"
        used_ids.add(nid)
        ntype = str(node.get("type") or "").strip().lower()
        if ntype not in _TOPOLOGY_TYPES:
            ntype = _infer_device_type(str(node.get("label") or ""))
        layer = str(node.get("layer") or "").strip().lower()
        if layer not in _TOPOLOGY_LAYERS:
            layer = _TOPOLOGY_TYPE_LAYER.get(ntype, "access")
        nodes.append({
            "id": nid,
            "label": str(node.get("label") or "").strip() or f"{ntype}-{index}",
            "type": ntype,
            "ip": str(node.get("ip") or "").strip(),
            "mac": str(node.get("mac") or "").strip(),
            "layer": layer,
            "zone": str(node.get("zone") or "").strip() or "内网区",
            "device_model": str(node.get("device_model") or "").strip(),
            "description": str(node.get("description") or "").strip(),
        })

    # 2. 边：只保留 source/target 均存在的边，避免悬空连线
    node_ids = {n["id"] for n in nodes}
    edges: list[dict] = []
    for index, edge in enumerate(data.get("edges") or [], start=1):
        if not isinstance(edge, dict):
            continue
        src, tgt = str(edge.get("source") or "").strip(), str(edge.get("target") or "").strip()
        if src in node_ids and tgt in node_ids and src != tgt:
            edges.append({
                "id": str(edge.get("id") or "").strip() or f"e-{index}",
                "source": src,
                "target": tgt,
                "label": str(edge.get("label") or "").strip(),
                "bandwidth": str(edge.get("bandwidth") or "").strip(),
            })

    # 3. 安全域：校验 nodes 引用，过滤空域；整体缺失时按节点 zone 聚合
    raw_zones = data.get("security_zones")
    zones: list[dict] = []
    if isinstance(raw_zones, list) and raw_zones:
        for zone in raw_zones:
            if not isinstance(zone, dict):
                continue
            zname = str(zone.get("zone_name") or "").strip()
            if not zname:
                continue
            zids = [nid for nid in (zone.get("nodes") or []) if nid in node_ids]
            if not zids:
                zids = [n["id"] for n in nodes if n["zone"] == zname]
            if not zids:
                continue
            zones.append({
                "zone_name": zname,
                "nodes": zids,
                "description": str(zone.get("description") or "").strip(),
                "risk_level": str(zone.get("risk_level") or "中").strip() or "中",
            })
    if not zones:
        zone_map: dict[str, list[str]] = {}
        for n in nodes:
            zone_map.setdefault(n["zone"], []).append(n["id"])
        zones = [
            {"zone_name": name, "nodes": ids, "description": "", "risk_level": "中"}
            for name, ids in zone_map.items()
        ]

    # 4. 摘要：缺失时按统计信息兜底生成
    summary = str(data.get("summary") or "").strip()
    if not summary and nodes:
        type_names = {
            "switch": "交换机", "router": "路由器", "firewall": "防火墙",
            "server": "服务器", "database": "数据库",
        }
        type_counts: dict[str, int] = {}
        for n in nodes:
            type_counts[n["type"]] = type_counts.get(n["type"], 0) + 1
        parts = [f"{type_names[t]}{c}台" for t, c in type_counts.items()]
        summary = (
            f"该网络按核心-汇聚-接入三层架构部署，共 {len(nodes)} 台设备"
            f"（{'、'.join(parts)}），划分为 {len(zones)} 个安全域"
            f"（{'、'.join(z['zone_name'] for z in zones)}）。"
            "建议进一步核查各设备访问控制策略与边界防护配置。"
        )
    return {"nodes": nodes, "edges": edges, "security_zones": zones, "summary": summary}


@tool
def search_tender(keyword: str, max_results: int = 10) -> str:
    """搜索招标公告网站，根据关键词查找招标信息。返回标题、URL、发布日期、截止日期、摘要。keyword: 搜索关键词（如"等保三级 招标"）"""
    start = time.perf_counter()
    try:
        from tavily import TavilyClient

        client = TavilyClient(api_key=settings.tavily_api_key)
        response = client.search(
            query=keyword,
            max_results=max_results or settings.tavily_max_results,
            include_domains=settings.tender_sites,
        )
        results = []
        for item in response.get("results") or []:
            results.append({
                "title": item.get("title", ""),
                "url": item.get("url", ""),
                "published_date": item.get("published_date", ""),
                "score": item.get("score", 0),
                "content": (item.get("content") or "")[:200],
            })
        logger.info(
            "search_tender 完成 keyword=%s results=%d duration=%dms",
            keyword, len(results), round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(results, ensure_ascii=False)
    except Exception as exc:
        logger.warning("search_tender 调用失败: %s", exc)
        return json.dumps({"error": "搜索服务暂时不可用", "detail": str(exc)}, ensure_ascii=False)


def search_tender_local(keywords: str, db: Session | None = None) -> dict:
    """本地历史标书库检索（kb_type=2），返回相似段落供 Agent 参考风格与结构。

    结果结构：{"tenders": [{"doc_id", "title", "snippet", "score"}], "reference_doc_ids": [...]}。
    """
    start = time.perf_counter()
    snippets = get_rag_context(keywords, [2])

    # 补充标书标题（kb_documents）
    doc_ids = {s["doc_id"] for s in snippets if s.get("doc_id") is not None}
    titles: dict = {}
    if doc_ids:
        own_session = db is None
        if db is None:
            db = SessionLocal()
        try:
            for doc in db.scalars(select(KbDocument).where(KbDocument.id.in_(doc_ids))).all():
                titles[doc.id] = doc.doc_title
        finally:
            if own_session:
                db.close()

    tenders = [
        {
            "doc_id": s["doc_id"],
            "title": titles.get(s["doc_id"], "历史标书"),
            "snippet": s["text"],
            "score": round(s["score"], 4),
        }
        for s in snippets
    ]
    return {
        "tenders": tenders,
        "reference_doc_ids": sorted(doc_ids),
        "duration_ms": round((time.perf_counter() - start) * 1000),
    }


@tool
def download_tender(url: str) -> str:
    """下载招标文件。从招标公告页面提取附件下载链接并下载。url: 招标公告页面 URL。返回本地文件路径"""
    start = time.perf_counter()
    try:
        import httpx
        from bs4 import BeautifulSoup

        headers = {"User-Agent": _TENDER_UA}
        # 1. 请求公告页面（30s 超时，跟随重定向，UA 伪装浏览器）
        resp = httpx.get(url, headers=headers, timeout=30.0, follow_redirects=True)
        resp.raise_for_status()

        # 2. 解析 HTML 提取附件下载链接（.pdf / .doc / .docx / .zip）
        soup = BeautifulSoup(resp.text, "lxml")
        links: list[str] = []
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if not href or href.startswith(("javascript:", "#", "mailto:", "tel:")):
                continue
            # 忽略查询串后判断扩展名（部分站点在链接后带 ?t=xxx 签名）
            path_only = href.lower().split("?")[0].split("#")[0]
            if path_only.endswith(_TENDER_FILE_EXTS):
                links.append(urljoin(url, href))
        if not links:
            logger.warning("download_tender 未找到附件链接 url=%s", url)
            return json.dumps(
                {"error": "未找到附件下载链接", "detail": "公告页面未解析到 .pdf/.doc/.docx/.zip 附件"},
                ensure_ascii=False,
            )

        # 3. 下载第一个附件到 tenders 目录，文件名 {时间戳}_{原文件名} 防冲突
        os.makedirs(settings.tender_dir, exist_ok=True)
        file_url = links[0]
        file_name = unquote(Path(urlparse(file_url).path).name) or f"tender_{int(time.time())}.pdf"
        file_path = str(Path(settings.tender_dir) / f"{int(time.time() * 1000)}_{file_name}")
        with httpx.stream("GET", file_url, headers=headers, timeout=30.0, follow_redirects=True) as fr:
            fr.raise_for_status()
            with open(file_path, "wb") as f:
                for chunk in fr.iter_bytes(chunk_size=64 * 1024):
                    f.write(chunk)
        file_size = os.path.getsize(file_path)
        logger.info(
            "download_tender 完成 url=%s 附件数=%d 耗时=%dms 大小=%dB",
            url, len(links), round((time.perf_counter() - start) * 1000), file_size,
        )
        return json.dumps(
            {"file_path": file_path, "file_name": file_name, "file_size": file_size},
            ensure_ascii=False,
        )
    except Exception as exc:
        logger.warning("download_tender 调用失败 url=%s: %s", url, exc)
        return json.dumps({"error": "招标文件下载失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def parse_tender(file_path: str) -> str:
    """解析招标文件，提取关键信息。file_path: 招标文件本地路径（或 minio:// 存储路径）。返回结构化 JSON，包含预算金额、评分标准、资格要求、废标项"""
    start = time.perf_counter()
    try:
        # 上传的招标文件存储在 MinIO（storage_path 形如 minio://{bucket}/{object_key}），
        # 需先拉取到本地临时文件再解析（与 review_bidding 相同处理）
        if file_path.startswith("minio://"):
            local_path = _resolve_bidding_file(file_path)
            if local_path is None:
                return json.dumps(
                    {"error": "文件不存在", "detail": f"{file_path} 无法访问"},
                    ensure_ascii=False,
                )
            file_path = local_path
        p = Path(file_path)
        if not p.exists():
            return json.dumps({"error": "文件不存在", "detail": f"{file_path} 不存在"}, ensure_ascii=False)

        # 1. 按后缀分流提取全文
        ext = p.suffix.lower()
        if ext == ".pdf":
            import fitz  # PyMuPDF

            doc = fitz.open(str(p))
            try:
                text = "\n".join(page.get_text() for page in doc)
            finally:
                doc.close()
        elif ext == ".docx":
            from docx import Document

            doc = Document(str(p))
            text = "\n".join(para.text for para in doc.paragraphs)
        elif ext == ".doc":
            # python-docx 无法解析真正的旧版 .doc（OLE），尝试读取，失败则建议转 PDF
            try:
                from docx import Document

                doc = Document(str(p))
                text = "\n".join(para.text for para in doc.paragraphs)
            except Exception:
                logger.warning("parse_tender 旧版 .doc 解析失败 file=%s", file_path)
                return json.dumps(
                    {"error": "旧版 .doc 格式无法直接解析", "detail": "请将文件转换为 .pdf 或 .docx 后重新解析"},
                    ensure_ascii=False,
                )
        else:
            return json.dumps(
                {"error": "不支持的文件格式", "detail": f"仅支持 .pdf/.docx/.doc，收到扩展名 {ext or '无'}"},
                ensure_ascii=False,
            )

        text = text.strip()
        if not text:
            return json.dumps({"error": "文件内容为空", "detail": "未能从文件中提取到任何文本"}, ensure_ascii=False)

        # 2. 超长文本截断，避免超出模型上下文
        raw_len = len(text)
        if raw_len > settings.parse_max_text_length:
            text = text[: settings.parse_max_text_length]
            logger.info("parse_tender 文本超长截断 %d -> %d 字符", raw_len, settings.parse_max_text_length)

        # 3. 交给 DeepSeek 做结构化提取
        llm = get_deepseek()
        extract_prompt = (
            "你是招投标文件解析助手。根据招标文件全文，提取以下字段并以 JSON 返回（只输出 JSON，不要多余内容）：\n"
            "- project_name：项目名称（字符串）\n"
            "- budget：预算金额（字符串，如\"80万元\"；未提及填 null）\n"
            "- deadline：投标截止日期（字符串，如\"2026-08-20\"；未提及填 null）\n"
            "- qualification_requirements：资格要求列表（字符串数组）\n"
            "- scoring_items：评分表项列表，每项含 item（评分项名称）、max_score（满分，数字）、description（评分说明）\n"
            "- disqualification_items：废标项列表（字符串数组）\n\n"
            f"招标文件全文：\n{text}"
        )
        response = llm.invoke(
            [
                {"role": "system", "content": "你只输出合法 JSON，不输出任何多余文字。"},
                {"role": "user", "content": extract_prompt},
            ],
            response_format={"type": "json_object"},
        )
        raw = (response.content or "").strip()
        # 清理可能出现的 Markdown 代码块包裹
        if raw.startswith("```"):
            raw = raw.strip("`")
            if raw.lower().startswith("json"):
                raw = raw[4:]
            raw = raw.strip()
        parsed = json.loads(raw)

        logger.info(
            "parse_tender 完成 file=%s 文本长度=%d 提取耗时=%dms",
            file_path, raw_len, round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(parsed, ensure_ascii=False)
    except Exception as exc:
        logger.warning("parse_tender 调用失败 file=%s: %s", file_path, exc)
        return json.dumps({"error": "招标文件解析失败", "detail": str(exc)}, ensure_ascii=False)


def analyze_screenshot(screenshot_path: str, checklist_code: str, checklist_name: str) -> dict:
    """VLM 分析测评截图，识别与指定控制点相关的配置信息。

    DASHSCOPE_API_KEY 缺失或模型调用失败时，如实返回错误信息（不伪造识别结果）。
    """
    start = time.perf_counter()
    try:
        vlm = get_qwen_vl()
        text = vlm.invoke([
            {
                "role": "user",
                "content": [
                    {"image": screenshot_path},
                    {"text": f"请识别该截图中与测评控制点「{checklist_name}」相关的配置信息，提取关键字段与取值。"},
                ],
            }
        ])
        return {
            "checklist_code": checklist_code,
            "checklist_name": checklist_name,
            "recognized": text,
            "status": "success",
            "duration_ms": round((time.perf_counter() - start) * 1000),
        }
    except Exception as exc:
        logger.warning("analyze_screenshot 调用失败: %s", exc)
        return {
            "checklist_code": checklist_code,
            "checklist_name": checklist_name,
            "recognized": "",
            "status": "error",
            "message": str(exc),
            "duration_ms": round((time.perf_counter() - start) * 1000),
        }


@tool
def generate_checklist(system_level: str, system_name: str = "") -> str:
    """根据等保级别生成测评控制点清单和对应测评命令。
    system_level: 等保级别，如"二级"/"三级"/"四级"
    system_name: 被测系统名称（可选，用于个性化命令）
    返回结构化 JSON，包含控制点列表
    """
    start = time.perf_counter()
    # 惰性导入：prompt_layer 顶层依赖本模块（search_tender 等），避免循环导入
    from .prompt_layer import get_checklist_generation_prompt
    try:
        prompt = get_checklist_generation_prompt()
        llm = get_deepseek()
        response = llm.invoke(
            prompt.format_messages(system_level=system_level, system_name=system_name or "被测系统"),
            response_format={"type": "json_object"},
        )
        result = _parse_llm_json(response.content)
        checklist = result.get("checklist") if isinstance(result, dict) else None
        if not isinstance(checklist, list) or not checklist:
            return json.dumps(
                {"error": "测评清单生成失败", "detail": "模型未返回有效控制点列表"},
                ensure_ascii=False,
            )
        result["system_level"] = system_level
        result["system_name"] = system_name or "被测系统"
        result["total_count"] = len(checklist)
        logger.info(
            "generate_checklist 完成 level=%s 控制点=%d 耗时=%dms",
            system_level, len(checklist), round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(result, ensure_ascii=False)
    except Exception as exc:
        logger.warning("generate_checklist 调用失败 level=%s: %s", system_level, exc)
        return json.dumps({"error": "测评清单生成失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def generate_topology(asset_text: str) -> str:
    """根据用户上传的资产核查表内容生成网络拓扑图数据。
    asset_text: 资产核查表的文本内容（设备名称/类型/IP/区域等设备清单，可直接传文件提取的全文）
    返回 topology_data JSON，包含 nodes/edges/security_zones/summary
    """
    start = time.perf_counter()
    # 惰性导入：prompt_layer 顶层依赖本模块（search_tender 等），避免循环导入
    from .prompt_layer import get_topology_generation_prompt
    try:
        prompt = get_topology_generation_prompt()
        llm = get_deepseek()
        # 长文本截断，避免超出模型上下文（资产表字段多，留足余量）
        truncated = (asset_text or "")[:8000]
        response = llm.invoke(
            prompt.format_messages(asset_text=truncated or "（空）"),
            response_format={"type": "json_object"},
        )
        result = _parse_llm_json(response.content)
        # 兜底规范化：补全字段/过滤无效边/补齐安全域与摘要，避免前端渲染异常
        result = _normalize_topology(result)
        nodes = result.get("nodes")
        if not isinstance(nodes, list) or not nodes:
            return json.dumps(
                {"error": "拓扑生成失败", "detail": "模型未返回有效节点列表"},
                ensure_ascii=False,
            )
        logger.info(
            "generate_topology 完成 节点=%d 边=%d 安全域=%d 耗时=%dms",
            len(nodes), len(result.get("edges") or []), len(result.get("security_zones") or []),
            round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(result, ensure_ascii=False)
    except Exception as exc:
        logger.warning("generate_topology 调用失败: %s", exc)
        return json.dumps({"error": "拓扑生成失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def compare_records(vlm_result: str, human_record: str, checklist_code: str) -> str:
    """比对 VLM 分析结果与人工记录，判断一致性并计算置信度。
    vlm_result: analyze_screenshot 返回的 JSON 字符串
    human_record: 组员人工填写的结果记录
    checklist_code: 测评控制点编号
    返回比对结果 JSON，含 consistent, ai_result, human_result, diff_detail, confidence
    """
    start = time.perf_counter()
    # 惰性导入：prompt_layer 顶层依赖本模块（search_tender 等），避免循环导入
    from .prompt_layer import get_comparison_prompt
    try:
        # 1. 从企业本地库（kb_type=1）检索特殊情况说明，作为置信度调整依据（失败降级为空，不阻塞比对）
        kb_refs: list = []
        try:
            from .memory import get_assessment_context

            kb_refs = get_assessment_context(checklist_code)
        except Exception as exc:
            logger.warning("compare_records 知识库检索失败（降级为空）: %s", exc)

        # 2. LLM 比对
        prompt = get_comparison_prompt()
        llm = get_deepseek()
        response = llm.invoke(
            prompt.format_messages(
                checklist_code=checklist_code,
                vlm_result=vlm_result or "（无 VLM 识别结果）",
                human_record=human_record or "（无人工记录）",
                kb_references=json.dumps(kb_refs, ensure_ascii=False) if kb_refs else "（无）",
            ),
            response_format={"type": "json_object"},
        )
        comparison = _parse_llm_json(response.content)

        # 3. 后端兜底：confidence 必须为 0~1 数字；KB 引用附到结果
        try:
            confidence = float(comparison.get("confidence") or 0.5)
        except (TypeError, ValueError):
            confidence = 0.5
        comparison["confidence"] = round(min(max(confidence, 0.0), 1.0), 2)
        comparison["kb_references"] = kb_refs

        logger.info(
            "compare_records 完成 code=%s confidence=%s 耗时=%dms",
            checklist_code, comparison["confidence"], round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(comparison, ensure_ascii=False)
    except Exception as exc:
        logger.warning("compare_records 调用失败 code=%s: %s", checklist_code, exc)
        return json.dumps({"error": "比对失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def generate_bidding(tender_info: str, rag_context: str, section: str = "all") -> str:
    """根据招标信息和历史标书参考，生成投标文件的指定部分。
    tender_info: parse_tender 返回的招标文件解析结果（JSON 字符串）
    rag_context: RAG 检索到的历史标书片段
    section: 生成哪个部分，可选 all/company_profile/qualification/project_team/pricing
    返回结构化 JSON，包含生成的各部分内容
    """
    start = time.perf_counter()
    # 惰性导入：prompt_layer 顶层依赖本模块（search_tender 等），避免循环导入
    from .prompt_layer import get_bidding_generation_prompt
    try:
        tender = json.loads(tender_info) if isinstance(tender_info, str) else (tender_info or {})
        project_name = str(tender.get("project_name") or "未知项目")[:50]

        prompt = get_bidding_generation_prompt()
        llm = get_deepseek()
        response = llm.invoke(
            prompt.format_messages(
                tender_info=tender_info,
                rag_context=rag_context or "（无参考标书）",
                section=section,
            ),
            response_format={"type": "json_object"},
        )
        sections = _parse_llm_json(response.content)

        # section 指定单部分时，只保留对应字段
        allowed = {"company_profile", "qualification", "project_team", "pricing"}
        if section != "all" and section in allowed:
            sections = {section: sections.get(section, "")}

        # 兜底：pricing 不允许出现具体金额，缺省时标注待人工审核
        pricing = str(sections.get("pricing") or "").strip()
        if "待人工审核" not in pricing and "待人工核对" not in pricing:
            sections["pricing"] = "（待人工审核）"

        word_counts = {k: len(str(v)) for k, v in sections.items()}
        logger.info(
            "generate_bidding 完成 project=%s section=%s 各部分字数=%s 耗时=%dms",
            project_name, section, word_counts, round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(sections, ensure_ascii=False)
    except Exception as exc:
        logger.warning("generate_bidding 调用失败 project=%s section=%s: %s", tender_info[:30], section, exc)
        return json.dumps({"error": "标书生成失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def self_check(bidding_sections: str, tender_info: str) -> str:
    """对生成的投标文件进行自检。逐条检查废标项是否满足、格式是否符合要求、错别字。
    bidding_sections: generate_bidding 返回的投标内容（JSON 字符串）
    tender_info: parse_tender 返回的招标文件解析结果（JSON 字符串）
    返回自检结果 JSON
    """
    start = time.perf_counter()
    from .prompt_layer import get_self_check_prompt
    try:
        sections = json.loads(bidding_sections) if isinstance(bidding_sections, str) else (bidding_sections or {})
        tender = json.loads(tender_info) if isinstance(tender_info, str) else (tender_info or {})
        project_name = str(tender.get("project_name") or "未知项目")[:50]

        prompt = get_self_check_prompt()
        llm = get_deepseek()
        response = llm.invoke(
            prompt.format_messages(
                tender_rules=json.dumps(tender, ensure_ascii=False),
                bidding_content=json.dumps(sections, ensure_ascii=False),
            ),
            response_format={"type": "json_object"},
        )
        result = _parse_llm_json(response.content)

        # 汇总统计（summary 缺失时由后端兜底计算）
        counts = {"pass": 0, "warning": 0, "danger": 0}
        for dim in ("disqualification_check", "format_check", "typo_check"):
            for item in result.get(dim) or []:
                status = str(item.get("status") or "").lower()
                if status in counts:
                    counts[status] += 1
        total = sum(counts.values())
        if not result.get("summary"):
            result["summary"] = (
                f"共检查{total}项，通过{counts['pass']}项，警告{counts['warning']}项，危险{counts['danger']}项"
            )
        if not result.get("overall_status"):
            result["overall_status"] = "danger" if counts["danger"] else "warning" if counts["warning"] else "pass"

        logger.info(
            "self_check 完成 project=%s 检查项=%d pass=%d warning=%d danger=%d 耗时=%dms",
            project_name, total, counts["pass"], counts["warning"], counts["danger"],
            round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(result, ensure_ascii=False)
    except Exception as exc:
        logger.warning("self_check 调用失败: %s", exc)
        return json.dumps({"error": "废标自检失败", "detail": str(exc)}, ensure_ascii=False)


def _resolve_bidding_file(path: str) -> str | None:
    """把投标文件路径解析为本地文件：minio:// 路径从 MinIO 拉取到临时文件，本地路径直接返回。

    用户上传的投标文件存储在 MinIO（storage_path 形如 minio://{bucket}/{object_key}），
    review_bidding 需先把对象下载到临时文件才能用 fitz / python-docx 提取文本。
    """
    if path.startswith("minio://"):
        try:
            from ..services.minio_client import get_minio_client, object_key_of

            object_key = object_key_of(path)
            if not object_key:
                logger.warning("review_bidding 无法解析 MinIO 路径: %s", path)
                return None
            client = get_minio_client()
            response = client.get_object(settings.minio_bucket, object_key)
            try:
                content = response.read()
            finally:
                response.close()
                response.release_conn()
            import tempfile

            suffix = Path(object_key).suffix or ".pdf"
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(content)
                return tmp.name
        except Exception as exc:
            logger.warning("review_bidding 从 MinIO 拉取文件失败 path=%s: %s", path, exc)
            return None
    return path if Path(path).exists() else None


def _extract_doc_text(file_path: str) -> tuple[str | None, str | None]:
    """提取文档全文（与 parse_tender 相同的 .pdf/.docx/.doc 分流逻辑）。

    返回 (文本, 错误信息)：成功时错误为 None，失败时文本为 None。
    """
    p = Path(file_path)
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            import fitz  # PyMuPDF

            doc = fitz.open(str(p))
            try:
                text = "\n".join(page.get_text() for page in doc)
            finally:
                doc.close()
        elif ext == ".docx":
            from docx import Document

            doc = Document(str(p))
            text = "\n".join(para.text for para in doc.paragraphs)
        elif ext == ".doc":
            # python-docx 无法解析真正的旧版 .doc（OLE），尝试读取，失败则建议转 PDF
            try:
                from docx import Document

                doc = Document(str(p))
                text = "\n".join(para.text for para in doc.paragraphs)
            except Exception:
                logger.warning("review_bidding 旧版 .doc 解析失败 file=%s", file_path)
                return None, "旧版 .doc 格式无法直接解析，请将文件转换为 .pdf 或 .docx 后重新审核"
        else:
            return None, f"仅支持 .pdf/.docx/.doc，收到扩展名 {ext or '无'}"
    except Exception as exc:
        logger.warning("review_bidding 文本提取失败 file=%s: %s", file_path, exc)
        return None, str(exc)

    text = (text or "").strip()
    if not text:
        return None, "未能从文件中提取到任何文本"
    if len(text) > settings.parse_max_text_length:
        text = text[: settings.parse_max_text_length]
    return text, None


@tool
def review_bidding(bidding_file_path: str, tender_info: str = "") -> str:
    """审核投标文件。检查废标项、格式、错字，返回审核结果。
    bidding_file_path: 投标文件本地路径（或 minio:// 存储路径）
    tender_info: 招标文件解析结果（JSON 字符串），可选。有则交叉比对，无则只查格式和错字
    返回审核结果 JSON
    """
    start = time.perf_counter()
    # 惰性导入：prompt_layer 顶层依赖本模块（search_tender 等），避免循环导入
    from .prompt_layer import get_bidding_review_prompt
    try:
        # 1. 定位文件并提取全文（复用 parse_tender 的分流逻辑）
        local_path = _resolve_bidding_file(bidding_file_path)
        if local_path is None:
            return json.dumps(
                {"error": "文件不存在", "detail": f"{bidding_file_path} 无法访问"},
                ensure_ascii=False,
            )
        text, err = _extract_doc_text(local_path)
        if text is None:
            return json.dumps({"error": "投标文件解析失败", "detail": err}, ensure_ascii=False)

        # 2. 交给 DeepSeek 做结构化审核（有 tender_info 交叉比对，无则通用检查）
        prompt = get_bidding_review_prompt(max_items=settings.bidding_review_max_items)
        llm = get_deepseek()
        response = llm.invoke(
            prompt.format_messages(
                tender_info=tender_info or "（无招标文件解析结果，做通用废标项检查）",
                bidding_content=text,
            ),
            response_format={"type": "json_object"},
        )
        result = _parse_llm_json(response.content)

        # 3. 后端兜底：need_human 由 confidence < 阈值决定；items 截断到最大检查项数
        items = result.get("items")
        if not isinstance(items, list):
            items = []
        items = items[: settings.bidding_review_max_items]
        counts = {"pass": 0, "warning": 0, "danger": 0, "human": 0}
        for item in items:
            if not isinstance(item, dict):
                continue
            try:
                conf = float(item.get("confidence") or 0.0)
            except (TypeError, ValueError):
                conf = 0.0
            item["confidence"] = round(conf, 2)
            if conf < settings.review_confidence_threshold:
                item["need_human"] = True
            status = str(item.get("status") or "").lower()
            if status in counts:
                counts[status] += 1
            if item.get("need_human"):
                counts["human"] += 1
        result["items"] = items
        if not result.get("overall_status"):
            result["overall_status"] = (
                "danger" if counts["danger"] else "warning" if counts["warning"] else "pass"
            )
        if not result.get("summary"):
            result["summary"] = (
                f"共检查{len(items)}项，通过{counts['pass']}项，警告{counts['warning']}项，"
                f"危险{counts['danger']}项，需人工确认{counts['human']}项"
            )

        logger.info(
            "review_bidding 完成 文本长度=%d 检查项=%d overall=%s 需人工=%d 耗时=%dms",
            len(text), len(items), result.get("overall_status"), counts["human"],
            round((time.perf_counter() - start) * 1000),
        )
        return json.dumps(result, ensure_ascii=False)
    except Exception as exc:
        logger.warning("review_bidding 调用失败: %s", exc)
        return json.dumps({"error": "投标文件审核失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def get_assessment_status(session_id: int) -> str:
    """查询当前测评会话的控制点状态汇总，用于回答"还有几个控制点没过/哪些待复核"等进度问题。
    session_id: 当前测评会话 ID（见系统提示中的"当前会话 ID"）
    返回 JSON：各状态数量 + 低置信度待复核控制点列表（含 record_id/编号/名称/置信度）
    """
    from ..db.session import SessionLocal
    from ..services.assessment_service import ensure_review_task  # noqa: F401 仅确保模块可导入

    try:
        with SessionLocal() as db:
            records = db.scalars(
                select(AssessmentRecord).where(AssessmentRecord.session_id == session_id)
            ).all()
            if not records:
                return json.dumps(
                    {"error": "该会话暂无测评控制点", "detail": "请先触发测评清单生成"},
                    ensure_ascii=False,
                )
            status_names = {0: "待测评", 2: "核查中", 3: "自动通过", 4: "待人工复核", 5: "已确认", 6: "异常"}
            counts = {k: 0 for k in status_names}
            review_items: list[dict] = []
            for r in records:
                counts[r.status] = counts.get(r.status, 0) + 1
                if r.status == 4:
                    review_items.append({
                        "record_id": r.id,
                        "checklist_code": r.checklist_code,
                        "checklist_name": r.checklist_name,
                        "confidence": r.confidence,
                    })
            summary = {
                "total": len(records),
                "counts": counts,
                "status_names": status_names,
                "review_items": review_items,
            }
            logger.info(
                "get_assessment_status 完成 session=%s total=%d review=%d",
                session_id, len(records), len(review_items),
            )
            return json.dumps(summary, ensure_ascii=False)
    except Exception as exc:
        logger.warning("get_assessment_status 调用失败 session=%s: %s", session_id, exc)
        return json.dumps({"error": "测评状态查询失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def confirm_assessment(record_id: int, approved: bool, session_id: int) -> str:
    """组长在对话中确认或驳回某个测评控制点的比对结果（低置信度复核）。
    record_id: 控制点记录 ID（可先调用 get_assessment_status 获取）
    approved: True 表示确认通过（status=5），False 表示驳回待整改（status=4 并保留复核任务）
    session_id: 当前测评会话 ID（见系统提示中的"当前会话 ID"）
    返回 JSON：record_id/checklist_code/checklist_name/status/confidence/确认结果说明
    """
    from ..db.session import SessionLocal
    from ..services.assessment_service import ensure_review_task, write_report_message

    try:
        with SessionLocal() as db:
            record = db.get(AssessmentRecord, record_id)
            if record is None or record.session_id != session_id:
                return json.dumps(
                    {"error": "测评记录不存在", "detail": "记录不存在或不属于当前会话"},
                    ensure_ascii=False,
                )
            now = int(time.time())
            record.updated_at = now
            if approved:
                record.status = 5
                # 完成该记录待审的复核任务（低置信度复核 task_type=1）
                pending = db.scalar(
                    select(ReviewTask).where(
                        ReviewTask.source_type == 1,
                        ReviewTask.source_id == record.id,
                        ReviewTask.task_type == 1,
                        ReviewTask.status == 0,
                    )
                )
                if pending is not None:
                    pending.review_result = 0
                    pending.status = 1
                    pending.reviewed_at = now
                action = f"已确认通过，置信度 {record.confidence:.0%}"
            else:
                record.status = 4
                ensure_review_task(db, record, now)
                action = f"已驳回待整改，置信度 {record.confidence:.0%}"
            write_report_message(
                db, record.user_id, record.session_id,
                f"【测评核查】第 {record.checklist_code} 项（{record.checklist_name}）{action}。",
            )
            db.commit()
            logger.info(
                "confirm_assessment 完成 record=%s approved=%s status=%s",
                record_id, approved, record.status,
            )
            return json.dumps({
                "record_id": record.id,
                "checklist_code": record.checklist_code,
                "checklist_name": record.checklist_name,
                "status": record.status,
                "confidence": record.confidence,
                "result": action,
            }, ensure_ascii=False)
    except Exception as exc:
        logger.warning("confirm_assessment 调用失败 record=%s: %s", record_id, exc)
        return json.dumps({"error": "确认操作失败", "detail": str(exc)}, ensure_ascii=False)


@tool
def submit_assessment_record(record_id: int, human_record: str, session_id: int) -> str:
    """在对话中为某个测评控制点提交人工测评记录；若该控制点已有 VLM 截图分析，则自动比对并更新状态。
    record_id: 控制点记录 ID（可先调用 get_assessment_status 获取）
    human_record: 组员实际填写的结果记录（如"锁定阈值为5"）
    session_id: 当前测评会话 ID（见系统提示中的"当前会话 ID"）
    返回 JSON：record_id/checklist_code/status/confidence/comparison_done
    """
    from ..db.session import SessionLocal
    from ..services.assessment_service import compare_and_update_record, write_report_message

    try:
        with SessionLocal() as db:
            record = db.get(AssessmentRecord, record_id)
            if record is None or record.session_id != session_id:
                return json.dumps(
                    {"error": "测评记录不存在", "detail": "记录不存在或不属于当前会话"},
                    ensure_ascii=False,
                )
            record.human_record = human_record
            comparison_done = False
            if record.vlm_analysis:
                compare_and_update_record(db, record)
                comparison_done = True
            record.updated_at = int(time.time())
            if comparison_done:
                if record.status == 3:
                    result_note = f"比对置信度 {record.confidence:.0%}，自动通过"
                elif record.status == 4:
                    result_note = f"比对置信度 {record.confidence:.0%}，待人工复核"
                else:
                    result_note = "比对完成"
            else:
                result_note = "尚未上传截图，待上传后自动比对"
            write_report_message(
                db, record.user_id, record.session_id,
                f"【测评核查】第 {record.checklist_code} 项（{record.checklist_name}）人工记录已提交，{result_note}。",
            )
            db.commit()
            logger.info(
                "submit_assessment_record 完成 record=%s compared=%s status=%s confidence=%s",
                record_id, comparison_done, record.status, record.confidence,
            )
            return json.dumps({
                "record_id": record.id,
                "checklist_code": record.checklist_code,
                "checklist_name": record.checklist_name,
                "status": record.status,
                "confidence": record.confidence,
                "comparison_done": comparison_done,
            }, ensure_ascii=False)
    except Exception as exc:
        logger.warning("submit_assessment_record 调用失败 record=%s: %s", record_id, exc)
        return json.dumps({"error": "人工记录提交失败", "detail": str(exc)}, ensure_ascii=False)
